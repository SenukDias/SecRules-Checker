from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from weasyprint import HTML

from app.models.db_models import Finding, Job

TEMPLATE_DIR = Path(__file__).parent / "templates"
_env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)), autoescape=select_autoescape(["html"]))

SEVERITY_FILL = {
    "critical": "D32F2F",
    "high": "E8620C",
    "medium": "B8960C",
    "low": "1F9D55",
    "info": "607D8B",
}


def _severity_counts(findings: list[Finding]) -> dict[str, int]:
    counts = Counter(f.severity for f in findings)
    for sev in ("critical", "high", "medium", "low", "info"):
        counts.setdefault(sev, 0)
    ordered = {sev: counts[sev] for sev in ("critical", "high", "medium", "low", "info")}
    return ordered


def _devices(job: Job) -> list[dict]:
    """Every device parsed from the uploaded file, as stored on the job (parsed_model)."""
    return (job.parsed_model or {}).get("devices", [])


def _interface_rows(devices: list[dict]) -> list[dict]:
    rows = []
    for device in devices:
        for iface in device.get("interfaces", []):
            rows.append({"device": device["name"], **iface})
    return rows


def _rule_rows(devices: list[dict]) -> list[dict]:
    rows = []
    for device in devices:
        for rule in device.get("rules", []):
            rows.append({"device": device["name"], **rule})
    return rows


def render_html(job: Job, findings: list[Finding], topology_image_b64: str | None = None) -> str:
    devices = _devices(job)
    template = _env.get_template("report.html")
    return template.render(
        job=job,
        findings=findings,
        severity_counts=_severity_counts(findings),
        generated_at=datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
        topology_image=topology_image_b64,
        devices=devices,
        interfaces=_interface_rows(devices),
        rules=_rule_rows(devices),
    )


def render_pdf(job: Job, findings: list[Finding], topology_image_b64: str | None = None) -> bytes:
    html_content = render_html(job, findings, topology_image_b64)
    return HTML(string=html_content, base_url=str(TEMPLATE_DIR)).write_pdf()


def render_excel(job: Job, findings: list[Finding]) -> bytes:
    from io import BytesIO

    devices = _devices(job)

    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    headers = ["Severity", "Category", "Device", "Rule", "Description", "Remediation"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2430", end_color="1F2430", fill_type="solid")

    for f in findings:
        ws.append([f.severity.upper(), f.category, f.device_name, f.rule_ref, f.description, f.remediation])
        row = ws[ws.max_row]
        fill_color = SEVERITY_FILL.get(f.severity, "607D8B")
        row[0].fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        row[0].font = Font(color="FFFFFF", bold=True)

    for col, width in zip("ABCDEF", (12, 22, 20, 20, 60, 60)):
        ws.column_dimensions[col].width = width

    _add_sheet(
        wb,
        "Devices",
        ["Name", "Vendor", "Type"],
        [[d.get("name"), d.get("vendor"), d.get("device_type")] for d in devices],
        widths=(24, 18, 14),
    )
    _add_sheet(
        wb,
        "Interfaces",
        ["Device", "Interface", "Zone", "IP Address", "Subnet Mask"],
        [
            [i.get("device"), i.get("name"), i.get("zone"), i.get("ip_address"), i.get("subnet_mask")]
            for i in _interface_rows(devices)
        ],
        widths=(20, 18, 18, 16, 16),
    )
    _add_sheet(
        wb,
        "Rule Base",
        ["Device", "Reference", "Position", "Name", "Action", "Source", "Destination", "Service", "Logging", "Disabled", "Description"],
        [
            [
                r.get("device"),
                r.get("id"),
                r.get("position"),
                r.get("name"),
                r.get("action"),
                ", ".join(r.get("source") or []),
                ", ".join(r.get("destination") or []),
                ", ".join(r.get("services") or []),
                r.get("logging_enabled"),
                r.get("disabled"),
                r.get("description"),
            ]
            for r in _rule_rows(devices)
        ],
        widths=(20, 24, 10, 20, 10, 24, 24, 20, 10, 10, 40),
    )

    raw_ws = wb.create_sheet("Raw Configuration")
    raw_ws.append(["Device", "Raw Uploaded Configuration"])
    for cell in raw_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2430", end_color="1F2430", fill_type="solid")
    for device in devices:
        raw_ws.append([device.get("name"), device.get("raw_source", "")])
        raw_ws[raw_ws.max_row][1].alignment = raw_ws[raw_ws.max_row][1].alignment.copy(wrap_text=True)
    raw_ws.column_dimensions["A"].width = 20
    raw_ws.column_dimensions["B"].width = 120

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: tuple[int, ...]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2430", end_color="1F2430", fill_type="solid")
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = width
